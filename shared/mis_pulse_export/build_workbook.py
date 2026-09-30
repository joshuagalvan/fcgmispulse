"""Pure function: rows in, an exact-format openpyxl Workbook out.

No FastAPI/SQLAlchemy/Qt dependency, so it can be built and inspected in
isolation (see backend/tests/test_export_format.py) against the real
historical file.
"""

from __future__ import annotations

import datetime as dt
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .recalc import recalc
from .workbook_template import (
    CENTERED_WRAPPED,
    COLUMNS,
    DATA_FONT,
    DATA_VALIDATIONS,
    DATE_FORMAT,
    DURATION_FORMAT,
    FREEZE_PANES,
    HEADER_FONT,
    THIN_BOX,
    TIME_FORMAT,
)


@dataclass
class EntryRow:
    entry_date: dt.date
    store_department: str
    reported_by: str
    time_sent: dt.time
    time_received: dt.time
    time_done: dt.time | None
    mis_personnel: str
    problem: str
    task: str
    type_: str
    type_of_support: str
    findings_cause: str | None = None
    action_taken: str | None = None
    remarks: str | None = None
    task_remarks: str | None = None
    acknowledged_by: str | None = None
    area_manager_head: str | None = None


def build_monthly_workbook(
    rows: list[EntryRow], month_name: str, year: int, *, validation_row_buffer: int = 2000
) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = f"{month_name.upper()} {year}"

    for idx, (letter, header, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=idx, value=header)
        cell.font = HEADER_FONT
        cell.alignment = CENTERED_WRAPPED
        cell.border = THIN_BOX
        if width is not None:
            ws.column_dimensions[letter].width = width

    for r, row in enumerate(rows, start=2):
        # Duration is only meaningful once Time Done is filled in — a blank
        # Time Done (an in-progress ticket) would otherwise evaluate the
        # formula against an implicit midnight and show a nonsense negative
        # duration, so leave it blank instead until Time Done is known.
        duration_value = f"=F{r}-E{r}" if row.time_done is not None else None
        values = [
            row.entry_date,
            row.store_department,
            row.reported_by,
            row.time_sent,
            row.time_received,
            row.time_done,
            duration_value,
            row.mis_personnel,
            row.problem,
            row.findings_cause,
            row.action_taken,
            row.remarks,
            row.task,
            row.task_remarks,
            row.acknowledged_by,
            row.area_manager_head,
            row.type_,
            row.type_of_support,
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col_idx, value=value)
            cell.font = DATA_FONT
            cell.alignment = CENTERED_WRAPPED
            cell.border = THIN_BOX

        ws.cell(row=r, column=1).number_format = DATE_FORMAT
        for col_idx in (4, 5, 6):
            ws.cell(row=r, column=col_idx).number_format = TIME_FORMAT
        ws.cell(row=r, column=7).number_format = DURATION_FORMAT

    ws.freeze_panes = FREEZE_PANES

    last_validated_row = max(len(rows) + 1, validation_row_buffer)
    for letter, choices in DATA_VALIDATIONS.items():
        dv = DataValidation(
            type="list", formula1='"' + ",".join(choices) + '"', allow_blank=True
        )
        ws.add_data_validation(dv)
        dv.add(f"{letter}2:{letter}{last_validated_row}")

    return wb


def save_and_recalc(wb: Workbook, path: str | Path, timeout: int = 90) -> dict:
    """Save the workbook and run it through LibreOffice so the Duration
    formulas have cached values, exactly like the skill's mandatory recalc
    step. Returns the recalc.py result dict; raises if it reports an error."""
    wb.save(str(path))
    result = recalc(str(path), timeout=timeout)
    if "error" in result:
        raise RuntimeError(f"Export recalculation failed: {result['error']}")
    return result
