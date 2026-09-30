"""Imports the real September 2026 data from the original spreadsheet.

A row counts as a real ticket only if it has both a valid Date (col A) and a
non-blank Problem (col I) -- analysis of the source file showed hundreds of
rows with a stray value in one column (e.g. a leftover Task dropdown pick or
a bare store name) but nothing else: leftover template rows, not tickets.

Store/brand lookups come from the authoritative Store List.xlsx via
seed_stores.py (run that first). This script still placeholder-seeds
`department`-type locations and `area_manager` from this data, since no
authoritative list for those exists yet -- and grows the `reported_by`
autocomplete pool.

Usage:
    ./.venv/bin/python scripts/seed_from_september.py "/path/to/MIS SUPPORT SUMMARY REPORT.xlsx"
"""

import datetime
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import load_workbook  # noqa: E402

from app import models  # noqa: E402
from app.db import Base, SessionLocal, engine  # noqa: E402
from app.services.brand_classify import classify  # noqa: E402

SHEET_NAME = "SEPTEMBER 2026"


def _clean(value):
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def read_real_rows(xlsx_path: str) -> list[dict]:
    wb = load_workbook(xlsx_path, data_only=False)
    ws = wb[SHEET_NAME]

    rows = []
    for r in range(2, ws.max_row + 1):
        entry_date = ws.cell(row=r, column=1).value
        problem = _clean(ws.cell(row=r, column=9).value)
        if not isinstance(entry_date, datetime.datetime) or not problem:
            continue

        def cell(col):
            return ws.cell(row=r, column=col).value

        rows.append(
            {
                "entry_date": entry_date.date(),
                "store_department": _clean(cell(2)) or "UNKNOWN",
                "reported_by": _clean(cell(3)) or "N/A",
                "time_sent": cell(4) if isinstance(cell(4), datetime.time) else None,
                "time_received": cell(5) if isinstance(cell(5), datetime.time) else None,
                "time_done": cell(6) if isinstance(cell(6), datetime.time) else None,
                "mis_personnel": _clean(cell(8)) or "",
                "problem": problem,
                "findings_cause": _clean(cell(10)),
                "action_taken": _clean(cell(11)),
                "remarks": _clean(cell(12)),
                "task": _clean(cell(13)) or "",
                "task_remarks": _clean(cell(14)),
                "acknowledged_by": _clean(cell(15)),
                "area_manager_head": _clean(cell(16)),
                "type_": _clean(cell(17)),
                "type_of_support": _clean(cell(18)),
            }
        )
    return rows


def main():
    xlsx_path = (
        sys.argv[1]
        if len(sys.argv) > 1
        else "/Users/joshuagalvan/Downloads/MIS SUPPORT SUMMARY REPORT.xlsx"
    )

    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        if db.query(models.Entry).count() > 0:
            print("entries table is not empty; skipping import to avoid duplicates.")
            return

        rows = read_real_rows(xlsx_path)

        users_by_name = {u.username.upper(): u.id for u in db.query(models.User).all()}
        admin = db.query(models.User).filter(models.User.is_admin.is_(True)).first()
        fallback_user_id = admin.id if admin else next(iter(users_by_name.values()))

        seen_departments: set[str] = set()
        seen_area_managers: set[str] = set()
        seen_reported_by: set[str] = set()

        for row in rows:
            personnel_key = row["mis_personnel"].strip().upper()
            created_by = users_by_name.get(personnel_key, fallback_user_id)
            loc_type, brand = classify(row["store_department"])

            db.add(
                models.Entry(
                    entry_date=row["entry_date"],
                    location_type=loc_type,
                    brand=brand,
                    store_department=row["store_department"],
                    reported_by=row["reported_by"],
                    time_sent=row["time_sent"],
                    time_received=row["time_received"],
                    time_done=row["time_done"],
                    mis_personnel=row["mis_personnel"],
                    problem=row["problem"],
                    findings_cause=row["findings_cause"],
                    action_taken=row["action_taken"],
                    remarks=row["remarks"],
                    task=row["task"],
                    task_remarks=row["task_remarks"],
                    acknowledged_by=row["acknowledged_by"],
                    area_manager_head=row["area_manager_head"],
                    type_=row["type_"],
                    type_of_support=row["type_of_support"],
                    created_by=created_by,
                )
            )

            if loc_type == "DEPARTMENT":
                seen_departments.add(row["store_department"])
            if row["area_manager_head"]:
                seen_area_managers.add(row["area_manager_head"])
            seen_reported_by.add(row["reported_by"])

        for value in seen_departments:
            db.add(models.Lookup(kind="location", location_type="DEPARTMENT", value=value))
        for value in seen_area_managers:
            db.add(models.Lookup(kind="area_manager", location_type=None, value=value))
        for value in seen_reported_by:
            db.add(models.Lookup(kind="reported_by", location_type=None, value=value))

        db.commit()

        print(f"Imported {len(rows)} entries from {SHEET_NAME}.")
        print(f"Seeded {len(seen_departments)} department lookups (placeholder, from this data).")
        print(f"Seeded {len(seen_area_managers)} area manager lookups (placeholder).")
        print(f"Seeded {len(seen_reported_by)} reported-by autocomplete values.")
    finally:
        db.close()


if __name__ == "__main__":
    main()
