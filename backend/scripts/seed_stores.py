"""Imports the authoritative store list (one tab per brand) as curated,
admin-managed lookups. Safe to re-run: existing (kind, brand, value) rows are
left alone, and brand-new ones are added.

Usage:
    ./.venv/bin/python scripts/seed_stores.py "/path/to/Store List.xlsx"
"""

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app import models  # noqa: E402
from app.db import Base, SessionLocal, engine  # noqa: E402


def _read_from_xlsx(xlsx_path: str) -> dict[str, list[str]]:
    from openpyxl import load_workbook

    wb = load_workbook(xlsx_path, data_only=True)
    by_brand = {}
    for brand in wb.sheetnames:
        ws = wb[brand]
        names = []
        for row in range(1, ws.max_row + 1):
            raw = ws.cell(row=row, column=1).value
            if raw is not None and str(raw).strip():
                names.append(" ".join(str(raw).split()))
        by_brand[brand] = names
    return by_brand


def main():
    xlsx_path = (
        sys.argv[1] if len(sys.argv) > 1 else "/Users/joshuagalvan/Downloads/Store List.xlsx"
    )

    if Path(xlsx_path).exists():
        by_brand = _read_from_xlsx(xlsx_path)
        print(f"Read store list from {xlsx_path}")
    else:
        from store_data import STORES_BY_BRAND as by_brand

        print(f"{xlsx_path} not found; using the embedded snapshot from store_data.py instead")

    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        brands_added = 0
        stores_added = 0
        stores_skipped_dupe = 0

        for brand, names in by_brand.items():
            existing_brand = (
                db.query(models.Lookup)
                .filter(models.Lookup.kind == "brand", models.Lookup.value == brand)
                .first()
            )
            if existing_brand is None:
                db.add(models.Lookup(kind="brand", value=brand))
                brands_added += 1

            seen_in_sheet = set()
            for name in names:
                if name in seen_in_sheet:
                    stores_skipped_dupe += 1
                    continue
                seen_in_sheet.add(name)

                existing_store = (
                    db.query(models.Lookup)
                    .filter(
                        models.Lookup.kind == "location",
                        models.Lookup.location_type == "STORE",
                        models.Lookup.brand == brand,
                        models.Lookup.value == name,
                    )
                    .first()
                )
                if existing_store is None:
                    db.add(
                        models.Lookup(
                            kind="location", location_type="STORE", brand=brand, value=name
                        )
                    )
                    stores_added += 1

        db.commit()
        print(f"Brands added: {brands_added}")
        print(f"Stores added: {stores_added}")
        print(f"Duplicate rows within a brand sheet skipped: {stores_skipped_dupe}")
    finally:
        db.close()


if __name__ == "__main__":
    main()
